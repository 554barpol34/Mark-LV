"""
Screenshot → Excel. JARVIS looks at the page the way a person does: it takes
screenshots, reads the fields the user asked for off the image, and adds them
to a workbook as rows. No select-all-and-copy, so it works on sites that block
copying or that load their text in ways the page text does not show.

The screenshots are kept next to the workbook (a folder with the same name),
so every row can be checked against what was on screen.
"""
from __future__ import annotations

import io
import json
import time
from datetime import datetime
from pathlib import Path

from jarvis import llm
from jarvis.tools import tool
from jarvis.tools.files import _shortcut

MAX_SHOTS = 6


def _browser_shots(screens: int) -> tuple[list[bytes], str]:
    from jarvis.tools.browser import BROWSER
    page = BROWSER._connect()
    shots = []
    for i in range(screens):
        shots.append(page.screenshot(type="jpeg", quality=80))
        if i < screens - 1:
            at_end = page.evaluate("innerHeight + scrollY >= document.documentElement.scrollHeight - 4")
            if at_end:
                break
            page.mouse.wheel(0, int(page.evaluate("innerHeight") * 0.85))
            time.sleep(0.7)
    return shots, page.url


def _screen_shot() -> list[bytes]:
    import mss
    from PIL import Image
    with mss.mss() as s:
        raw = s.grab(s.monitors[1])
    img = Image.frombytes("RGB", raw.size, raw.rgb)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=80)
    return [buf.getvalue()]


def _read(shots: list[bytes], fields: list[str], many: bool, notes: str) -> list[dict]:
    from google.genai import types
    shape = "one row per listing/item visible" if many else "exactly ONE row for the item this page is about"
    prompt = (
        "These are screenshots of one screen/page, top to bottom (they may overlap).\n"
        f"Read these fields: {json.dumps(fields, ensure_ascii=False)}.\n"
        f"Return JSON {{\"rows\": [{{field: value, ...}}]}} with {shape}. Use the field names "
        "exactly as given. Copy values exactly as shown (numbers, phone numbers, IDs digit for "
        "digit). If a field is not visible, use an empty string; never guess or invent."
        + (f"\nExtra instructions: {notes}" if notes else ""))
    parts = [types.Part.from_bytes(data=b, mime_type="image/jpeg") for b in shots]
    parts.append(types.Part.from_text(text=prompt))
    resp = llm.generate([types.Content(role="user", parts=parts)], json=True, thinking=1024)
    text = (resp.text or "").strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        data = json.loads(text[text.find("{"):text.rfind("}") + 1])
    rows = data.get("rows", data) if isinstance(data, dict) else data
    return [r for r in rows if isinstance(r, dict) and any(str(v).strip() for v in r.values())]


def _append(path: Path, fields: list[str], rows: list[dict], extra: dict) -> tuple[int, int]:
    """Append rows under matching headers; skip rows whose first field is already in the sheet."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from jarvis.tools.excel import _autosize, _value
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        header = [c.value for c in ws[1]]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Veriler"
        header = []
    for col in fields + list(extra):
        if col not in header:
            header.append(col)
            ws.cell(row=1, column=len(header), value=col)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"
    key_col = header.index(fields[0]) + 1
    seen = {str(ws.cell(row=r, column=key_col).value).strip() for r in range(2, ws.max_row + 1)}
    added = skipped = 0
    for row in rows:
        key = str(row.get(fields[0], "")).strip()
        if key and key in seen:
            skipped += 1
            continue
        seen.add(key)
        values = {**row, **extra}
        ws.append([_value(values.get(h, "")) if h in fields else values.get(h, "") for h in header])
        added += 1
    _autosize(ws)
    wb.save(path)
    return added, skipped


@tool("capture_to_excel", """Take screenshots, READ the requested fields off the images, and add
them as rows to an Excel file. Use this to collect data from listings and pages (ilan no, fiyat,
telefon, m², adres…), never select-all/copy. Typical loop: open a listing → (click "show phone"
if the user wants the number) → capture_to_excel → back → next listing. On a results page with
many items use many=true. The first field is the key: a row whose key is already in the file is
skipped, so running it twice on the same listing is safe.""",
      {"fields": {"type": "array", "items": {"type": "string"},
                  "description": "column names, in the user's words, e.g. [\"İlan No\", \"Fiyat\", \"Telefon\"]"},
       "path": {"type": "string", "description": "workbook, e.g. beykoz_riva_ilanlar.xlsx"},
       "source": {"type": "string", "enum": ["browser", "screen"],
                  "description": "browser = JARVIS's browser page (default); screen = whatever is on the monitor"},
       "screens": {"type": "integer", "description": "how many screens to capture while scrolling down (1-6, default 2)"},
       "many": {"type": "boolean", "description": "true = one row per item on the page (a results list)"},
       "notes": {"type": "string", "description": "extra reading rules from the user, optional"}},
      ["fields", "path"],
      label=lambda a: f"Ekran görüntüsünden okunuyor → {Path(a.get('path', '')).name}")
def capture_to_excel(fields: list, path: str, source: str = "browser", screens: int = 2,
                     many: bool = False, notes: str = ""):
    fields = [str(f).strip() for f in fields if str(f).strip()]
    if not fields:
        return "ERROR: give at least one field to read."
    p = _shortcut(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        p = p.with_suffix(".xlsx")
    screens = max(1, min(int(screens or 1), MAX_SHOTS))
    if source == "screen":
        shots, url = _screen_shot(), ""
    else:
        shots, url = _browser_shots(screens)

    folder = p.with_name(p.stem + "_ekran_goruntuleri")
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for i, b in enumerate(shots, 1):
        (folder / f"{stamp}_{i}.jpg").write_bytes(b)

    rows = _read(shots, fields, many, notes)
    if not rows:
        return "Could not read any of those fields from the screenshots. Check the page with browser_screenshot."
    extra = {"Kaynak": url} if url else {}
    extra["Ekran görüntüsü"] = f"{folder.name}/{stamp}_1.jpg"
    try:
        added, skipped = _append(p, fields, rows, extra)
    except PermissionError:
        return f"Could not save {p.name}: it is open in Excel. Ask the user to close it."
    preview = "\n".join(" | ".join(f"{k}: {r.get(k, '')}" for k in fields) for r in rows[:5])
    note = f", {skipped} already in the file" if skipped else ""
    return f"Read {len(rows)} row(s) from {len(shots)} screenshot(s); added {added}{note} to {p}.\n{preview}"
