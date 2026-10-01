"""
Excel files: create, fill, read, edit, chart. Works on .xlsx with openpyxl, so
it does not need Excel to be open; open_file shows the result in Excel.

If the file is open in Excel, Windows locks it and saving fails; the tool says
so instead of losing the edit.
"""
from __future__ import annotations

from pathlib import Path

from jarvis.tools import tool
from jarvis.tools.files import _shortcut


def _path(path: str) -> Path:
    p = _shortcut(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        p = p.with_suffix(".xlsx")
    return p


def _save(wb, p: Path) -> str | None:
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)
        return None
    except PermissionError:
        return (f"Could not save {p.name}: it is open in Excel. Ask the user to close it, "
                "or save under another name.")


def _value(v):
    """Numbers that arrive as text ("1.250", "42,5", "%12") become numbers."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if s.startswith("="):
        return s
    if (s[:1] == "0" and s[1:2].isdigit()) or len(s) > 15:
        return v                       # phone numbers, IDs: keep the leading zero
    t = s.replace("₺", "").replace("TL", "").replace("$", "").replace("€", "").replace("%", "").strip()
    if t and all(c.isdigit() or c in ".,-" for c in t) and any(c.isdigit() for c in t):
        if "," in t and "." in t:
            t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
        elif "," in t:
            t = t.replace(",", ".") if len(t.split(",")[-1]) != 3 else t.replace(",", "")
        elif t.count(".") > 1 or (t.count(".") == 1 and len(t.split(".")[-1]) == 3 and len(t) > 4):
            t = t.replace(".", "")
        try:
            n = float(t)
            return int(n) if n.is_integer() and "." not in t else n
        except ValueError:
            pass
    return v


def _autosize(ws):
    from openpyxl.utils import get_column_letter
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(8, w + 2), 60)


@tool("excel_write", """Write rows into an Excel file. rows is a list of rows, each a list of
cells; put the header row first. mode: "new" (default, replaces the sheet), "append" (adds rows
under existing data). Numbers given as text are stored as numbers; formulas start with "=".
Returns the full path; then call open_file if the user wants to see it.""",
      {"path": {"type": "string", "description": "e.g. arsalar.xlsx (saved in the workspace)"},
       "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}},
                "description": "cells as text; numbers like 4.250.000 TL become numbers"},
       "sheet": {"type": "string"},
       "mode": {"type": "string", "enum": ["new", "append"]}},
      ["path", "rows"], label=lambda a: f"Excel: {len(a.get('rows') or [])} satır → {Path(a.get('path', '')).name}")
def excel_write(path: str, rows: list, sheet: str = "", mode: str = "new"):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    p = _path(path)
    wb = load_workbook(p) if p.exists() else Workbook()
    if not p.exists() and wb.active.max_row == 1 and wb.active["A1"].value is None:
        wb.remove(wb.active)
    name = (sheet or "Sayfa1")[:31]
    if name in wb.sheetnames and mode == "new":
        idx = wb.sheetnames.index(name)
        wb.remove(wb[name])
        ws = wb.create_sheet(name, idx)
    elif name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
    first = ws.max_row + 1 if mode == "append" and ws.max_row > 1 else 1
    for r in rows:
        ws.append([_value(v) for v in (r if isinstance(r, list) else [r])])
    if first == 1 and rows:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    _autosize(ws)
    err = _save(wb, p)
    return err or f"Saved {len(rows)} rows to {p} (sheet {name})."


@tool("excel_read", "Read an Excel file (all sheets, or one) as tab-separated rows.",
      {"path": {"type": "string"}, "sheet": {"type": "string"}, "max_rows": {"type": "integer"}},
      ["path"], label=lambda a: f"Excel okunuyor: {Path(a.get('path', '')).name}")
def excel_read(path: str, sheet: str = "", max_rows: int = 300):
    from openpyxl import load_workbook
    p = _shortcut(path)
    wb = load_workbook(p, data_only=True, read_only=True)
    formulas = load_workbook(p, read_only=True)     # formulas Excel has not calculated yet
    out = []
    for ws in ([wb[sheet]] if sheet else wb.worksheets):
        out.append(f"== Sheet {ws.title} ==")
        raw = formulas[ws.title].iter_rows(values_only=True)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            f_row = next(raw, ())
            row = tuple(f_row[j] if v is None and j < len(f_row) else v for j, v in enumerate(row))
            if i >= max_rows:
                out.append(f"… more rows (max_rows={max_rows})")
                break
            if any(v is not None for v in row):
                out.append(f"{i + 1}\t" + "\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


@tool("excel_set_cells", """Set individual cells, e.g. [{"cell": "B2", "value": "150"},
{"cell": "C10", "value": "=SUM(C2:C9)"}]. Creates the file if needed.""",
      {"path": {"type": "string"}, "sheet": {"type": "string"},
       "cells": {"type": "array", "items": {"type": "object", "properties": {
           "cell": {"type": "string"}, "value": {"type": "string"}}, "required": ["cell", "value"]}}},
      ["path", "cells"], label=lambda a: f"Excel düzenleniyor: {Path(a.get('path', '')).name}")
def excel_set_cells(path: str, cells: list, sheet: str = ""):
    from openpyxl import Workbook, load_workbook
    p = _path(path)
    wb = load_workbook(p) if p.exists() else Workbook()
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else (wb.create_sheet(sheet[:31]) if sheet else wb.active)
    if isinstance(cells, dict):
        cells = [{"cell": k, "value": v} for k, v in cells.items()]
    for c in cells:
        ws[c["cell"].upper()] = _value(c["value"])
    _autosize(ws)
    return _save(wb, p) or f"Updated {len(cells)} cells in {p}."


@tool("excel_chart", """Add a chart to a sheet from its data. The first column is the labels, the
given value columns are the series; the header row names them.""",
      {"path": {"type": "string"}, "sheet": {"type": "string"},
       "kind": {"type": "string", "enum": ["bar", "line", "pie"]},
       "value_columns": {"type": "array", "items": {"type": "integer"},
                         "description": "1-based column numbers, e.g. [2] for column B"},
       "title": {"type": "string"}},
      ["path", "value_columns"], label=lambda a: "Excel: grafik ekleniyor")
def excel_chart(path: str, value_columns: list, sheet: str = "", kind: str = "bar", title: str = ""):
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    p = _path(path)
    wb = load_workbook(p)
    ws = wb[sheet] if sheet else wb.active
    chart = {"line": LineChart, "pie": PieChart}.get(kind, BarChart)()
    chart.title = title or None
    n = ws.max_row
    for col in value_columns:
        chart.add_data(Reference(ws, min_col=int(col), min_row=1, max_row=n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    chart.width, chart.height = 18, 9
    from openpyxl.utils import get_column_letter
    ws.add_chart(chart, f"{get_column_letter(ws.max_column + 2)}2")
    return _save(wb, p) or f"Added a {kind} chart to {p.name}."
